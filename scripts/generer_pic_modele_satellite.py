#!/usr/bin/env python3
"""PIC style modèle — fond satellite réel Pont Kiridi + zones numérotées."""

from pathlib import Path
import numpy as np
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
from matplotlib.patches import (
    FancyBboxPatch, Rectangle, Polygon, Circle, FancyArrowPatch,
    Ellipse, Arc, Wedge
)
from matplotlib.lines import Line2D
from PIL import Image as PILImage
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE

OUT = Path("/workspace/exports/plan-installation-chantier")
ART = Path("/opt/cursor/artifacts/plan-installation-chantier")
OUT.mkdir(parents=True, exist_ok=True)
ART.mkdir(parents=True, exist_ok=True)

BASE = OUT / "base_satellite_kiridi.png"
NAVY = "#0A2540"
RED = "#C0392B"
YELLOW = "#F4D03F"


def icon_building(ax, x, y, s=1.0, color="#5DADE2"):
    """Petit bâtiment modulaire 3D schématique."""
    # face
    ax.add_patch(Rectangle((x, y), 2.2 * s, 1.4 * s, facecolor=color, edgecolor="#1A5276", lw=0.8, zorder=20))
    # toit
    ax.add_patch(Polygon([
        (x, y + 1.4 * s), (x + 0.5 * s, y + 1.9 * s), (x + 2.7 * s, y + 1.9 * s), (x + 2.2 * s, y + 1.4 * s)
    ], closed=True, facecolor="#2874A6", edgecolor="#1A5276", lw=0.7, zorder=21))
    # côté
    ax.add_patch(Polygon([
        (x + 2.2 * s, y), (x + 2.7 * s, y + 0.5 * s), (x + 2.7 * s, y + 1.9 * s), (x + 2.2 * s, y + 1.4 * s)
    ], closed=True, facecolor="#3498DB", edgecolor="#1A5276", lw=0.7, zorder=20))
    # fenêtres
    for dx in [0.3, 0.9, 1.5]:
        ax.add_patch(Rectangle((x + dx * s, y + 0.45 * s), 0.35 * s, 0.4 * s,
                               facecolor="#FEF9E7", edgecolor="#1A5276", lw=0.4, zorder=22))


def icon_rebar(ax, x, y, s=1.0):
    for i, dy in enumerate([0, 0.35, 0.7, 1.05]):
        ax.add_patch(Rectangle((x, y + dy * s), 2.4 * s, 0.22 * s,
                               facecolor="#7F8C8D", edgecolor="#2C3E50", lw=0.5, zorder=20))
        ax.plot([x + 0.2 * s, x + 2.2 * s], [y + (dy + 0.11) * s] * 2,
                color="#BDC3C7", lw=0.6, zorder=21)


def icon_aggregates(ax, x, y, s=1.0):
    ax.add_patch(Polygon([
        (x, y), (x + 1.2 * s, y + 1.6 * s), (x + 2.6 * s, y),
    ], closed=True, facecolor="#A569BD", edgecolor="#6C3483", lw=0.8, zorder=20))
    ax.add_patch(Polygon([
        (x + 0.8 * s, y), (x + 2.0 * s, y + 1.1 * s), (x + 3.2 * s, y),
    ], closed=True, facecolor="#BB8FCE", edgecolor="#6C3483", lw=0.7, zorder=19, alpha=0.9))


def icon_beams(ax, x, y, s=1.0):
    for i, dy in enumerate([0, 0.45, 0.9]):
        ax.add_patch(FancyBboxPatch((x + i * 0.15 * s, y + dy * s), 2.8 * s, 0.35 * s,
                                    boxstyle="round,pad=0.01,rounding_size=0.05",
                                    facecolor="#85929E", edgecolor="#2C3E50", lw=0.7, zorder=20))
    # chariot / engins
    ax.add_patch(Rectangle((x + 3.0 * s, y), 1.2 * s, 0.7 * s, facecolor="#F39C12",
                           edgecolor="#9A7D0A", lw=0.6, zorder=20))
    ax.add_patch(Circle((x + 3.25 * s, y), 0.18 * s, facecolor="#2C3E50", zorder=21))
    ax.add_patch(Circle((x + 3.95 * s, y), 0.18 * s, facecolor="#2C3E50", zorder=21))


def icon_trucks(ax, x, y, s=1.0):
    for i, dx in enumerate([0, 2.6]):
        ax.add_patch(FancyBboxPatch((x + dx * s, y), 2.2 * s, 0.9 * s,
                                    boxstyle="round,pad=0.02,rounding_size=0.1",
                                    facecolor="#E67E22", edgecolor="#A04000", lw=0.7, zorder=20))
        ax.add_patch(Rectangle((x + (dx + 0.15) * s, y + 0.9 * s), 1.1 * s, 0.55 * s,
                               facecolor="#F5B041", edgecolor="#A04000", lw=0.5, zorder=21))
        ax.add_patch(Circle((x + (dx + 0.45) * s, y), 0.22 * s, facecolor="#2C3E50", zorder=22))
        ax.add_patch(Circle((x + (dx + 1.7) * s, y), 0.22 * s, facecolor="#2C3E50", zorder=22))


def icon_spoil(ax, x, y, s=1.0):
    ax.add_patch(Polygon([
        (x, y), (x + 1.4 * s, y + 1.5 * s), (x + 3.0 * s, y)
    ], closed=True, facecolor="#A04000", edgecolor="#6E2C00", lw=0.8, zorder=20))
    ax.add_patch(Polygon([
        (x + 0.9 * s, y), (x + 2.2 * s, y + 1.0 * s), (x + 3.6 * s, y)
    ], closed=True, facecolor="#CA6F1E", edgecolor="#6E2C00", lw=0.6, zorder=19, alpha=0.85))


def icon_fuel(ax, x, y, s=1.0):
    ax.add_patch(Ellipse((x + 0.9 * s, y + 0.7 * s), 1.8 * s, 1.4 * s,
                         facecolor="#E74C3C", edgecolor="#922B21", lw=0.8, zorder=20))
    ax.add_patch(Rectangle((x + 0.7 * s, y + 1.3 * s), 0.4 * s, 0.35 * s,
                           facecolor="#922B21", zorder=21))
    ax.text(x + 0.9 * s, y + 0.65 * s, "FUEL", fontsize=6 * s, ha="center", va="center",
            color="white", fontweight="bold", zorder=22)


def icon_generator(ax, x, y, s=1.0):
    ax.add_patch(FancyBboxPatch((x, y), 1.6 * s, 1.0 * s, boxstyle="round,pad=0.02",
                                facecolor="#566573", edgecolor="#1C2833", lw=0.7, zorder=20))
    ax.add_patch(Circle((x + 0.8 * s, y + 0.5 * s), 0.28 * s, facecolor="#F7DC6F",
                        edgecolor="#1C2833", lw=0.5, zorder=21))


def icon_assembly(ax, x, y, s=1.0):
    ax.add_patch(Circle((x, y), 1.3 * s, facecolor="#196F3D", edgecolor="white", lw=1.5, zorder=20))
    ax.plot([x - 0.7 * s, x, x + 0.7 * s, x, x - 0.7 * s],
            [y, y + 0.7 * s, y, y - 0.7 * s, y], color="white", lw=1.4, zorder=21)
    ax.plot(x, y, "o", color="white", ms=4, zorder=22)


def icon_barrier(ax, x, y, n=3):
    for i in range(n):
        xx = x + i * 0.55
        ax.add_patch(Rectangle((xx, y), 0.5, 0.35, facecolor="#E74C3C" if i % 2 == 0 else "white",
                               edgecolor="#922B21", lw=0.4, zorder=25))


def zone_box(ax, rect, color, number, title, details=None, alpha=0.35):
    x, y, w, h = rect
    ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.02,rounding_size=0.4",
                                facecolor=color, edgecolor=color, lw=2.4, alpha=alpha, zorder=10))
    # bordure pleine plus marquée
    ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.02,rounding_size=0.4",
                                fill=False, edgecolor=color, lw=2.6, zorder=11))
    # pastille numéro
    ax.add_patch(Circle((x + 0.9, y + h - 0.9), 0.75, facecolor=color, edgecolor="white",
                        lw=1.5, zorder=30))
    ax.text(x + 0.9, y + h - 0.9, str(number), color="white", ha="center", va="center",
            fontsize=11, fontweight="bold", zorder=31)
    ax.text(x + w / 2, y + h - 0.85, title, ha="center", va="center", fontsize=8,
            fontweight="bold", color="#1C2833", zorder=31,
            bbox=dict(boxstyle="round,pad=0.15", facecolor="white", edgecolor="none", alpha=0.85))
    if details:
        ax.text(x + w / 2, y + 1.1, details, ha="center", va="bottom", fontsize=6.5,
                color="#1C2833", zorder=31,
                bbox=dict(boxstyle="round,pad=0.12", facecolor="white", edgecolor="none", alpha=0.8))


def draw_arrow_path(ax, pts, color=YELLOW, lw=3.5):
    xs, ys = zip(*pts)
    ax.plot(xs, ys, color="#B7950B", lw=lw + 1.5, solid_capstyle="round", zorder=18)
    ax.plot(xs, ys, color=color, lw=lw, solid_capstyle="round", zorder=19)
    for i in range(len(pts) - 1):
        ax.annotate("", xy=pts[i + 1], xytext=pts[i],
                    arrowprops=dict(arrowstyle="->", color="#9A7D0A", lw=1.8),
                    zorder=20)


def fig_pic_modele():
    img = PILImage.open(BASE).convert("RGB")
    W, H = img.size  # 1280x1280

    fig, ax = plt.subplots(figsize=(16, 14), dpi=200)
    ax.imshow(img, extent=[0, 100, 0, 100], origin="upper")
    ax.set_xlim(0, 100)
    ax.set_ylim(0, 100)
    ax.set_aspect("equal")
    ax.axis("off")

    # --- Emprise / limite chantier (rouge tirets) ---
    # Adaptée au site : autour du franchissement + parcelles libres proches
    limite = [
        (48, 38), (72, 38), (78, 48), (82, 58), (78, 68),
        (70, 74), (52, 74), (44, 66), (42, 52), (48, 38),
    ]
    # Plusieurs îlots comme sur le modèle
    limite2 = [(28, 42), (42, 42), (42, 58), (28, 58)]
    limite3 = [(74, 30), (90, 30), (90, 46), (74, 46)]
    for poly in [limite, limite2, limite3]:
        xs, ys = zip(*poly)
        ax.plot(list(xs) + [xs[0]], list(ys) + [ys[0]], color=RED, lw=2.2, ls="--", zorder=12)

    # Zone pont en construction
    ax.add_patch(FancyBboxPatch((54, 52), 16, 8, boxstyle="round,pad=0.05,rounding_size=0.3",
                                facecolor="#E74C3C", edgecolor=RED, lw=2.5, alpha=0.28, zorder=9))
    ax.text(62, 56, "ZONE DE CHANTIER\nPONT EN CONSTRUCTION\nL = 62,80 m",
            ha="center", va="center", fontsize=7.5, fontweight="bold", color=RED, zorder=32,
            bbox=dict(boxstyle="round,pad=0.2", facecolor="white", edgecolor=RED, lw=1.2, alpha=0.92))

    # Sens d'écoulement (rivière ~ Ouest → Est sur l'image)
    ax.annotate("", xy=(78, 64), xytext=(48, 66),
                arrowprops=dict(arrowstyle="->", color="#2980B9", lw=2.8), zorder=15)
    ax.text(62, 68.5, "Sens d'écoulement", fontsize=8, color="#1A5276",
            fontweight="bold", ha="center",
            bbox=dict(boxstyle="round,pad=0.15", facecolor="white", alpha=0.85), zorder=32)

    # ========== ZONES 1–9 (style modèle) ==========
    # 1 Base vie — parcelle ouest du franchissement
    zone_box(ax, (28.5, 44, 12.5, 13), "#5DADE2", 1, "Base vie",
             "Bureau • Vestiaires\nSanitaires • Réfectoire", alpha=0.40)
    icon_building(ax, 31.5, 46.5, s=1.15)

    # 2 Magasin stockage
    zone_box(ax, (29, 30, 13, 10.5), "#F7DC6F", 2, "Magasin de stockage",
             "Ciment • Acier\nCoffrages • Bois", alpha=0.42)
    icon_building(ax, 32, 32, s=1.0, color="#F4D03F")

    # 3 Aire de ferraillage
    zone_box(ax, (44, 30, 12, 9), "#82E0AA", 3, "Aire de ferraillage",
             "Coupe / façonnage HA", alpha=0.40)
    icon_rebar(ax, 46.5, 32.2, s=1.1)

    # 4 Stockage agrégats
    zone_box(ax, (58, 28, 11.5, 9.5), "#BB8FCE", 4, "Stockage agrégats",
             "Sable • Graviers", alpha=0.40)
    icon_aggregates(ax, 60.5, 29.5, s=1.1)

    # 5 Aire de préfabrication
    zone_box(ax, (74.5, 31.5, 14.5, 13), "#5DADE2", 5, "Aire de préfabrication",
             "Poutres BA • Maturation", alpha=0.38)
    icon_beams(ax, 76.5, 34, s=1.05)

    # 6 Stationnement camions
    zone_box(ax, (74.5, 48, 14, 9.5), "#E67E22", 6, "Stationnement camions",
             "Livraisons / évacuation", alpha=0.38)
    icon_trucks(ax, 76.5, 50, s=0.95)

    # 7 Dépôt des déblais
    zone_box(ax, (44, 40.5, 10, 8.5), "#A04000", 7, "Dépôt des déblais",
             "Terrassements", alpha=0.40)
    icon_spoil(ax, 45.5, 41.5, s=1.0)

    # 8 Équipements annexes
    zone_box(ax, (18, 44, 9.5, 12), "#E67E22", 8, "Équipements annexes",
             "Cuve carburant\nGroupe • Atelier", alpha=0.38)
    icon_fuel(ax, 19.5, 49.5, s=0.85)
    icon_generator(ax, 22.5, 46, s=0.9)

    # 9 Point de rassemblement
    zone_box(ax, (48, 67, 10, 6.5), "#196F3D", 9, "Point de rassemblement",
             alpha=0.40)
    icon_assembly(ax, 53, 69.2, s=1.0)

    # --- Déviation provisoire ---
    detour = [
        (50, 48), (46, 36), (55, 24), (70, 22), (84, 28), (88, 42), (82, 52)
    ]
    draw_arrow_path(ax, detour, color=YELLOW, lw=3.2)
    ax.text(66, 19.5, "Déviation provisoire", fontsize=9, fontweight="bold",
            color="#9A7D0A", ha="center",
            bbox=dict(boxstyle="round,pad=0.2", facecolor="white", edgecolor=YELLOW, lw=1.5),
            zorder=33)

    # Accès chantier (flèches jaunes tirets)
    for (a, b) in [((42, 50), (48, 54)), ((78, 46), (70, 54)), ((55, 42), (58, 50))]:
        ax.annotate("", xy=b, xytext=a,
                    arrowprops=dict(arrowstyle="->", color="#F1C40F", lw=2.0, linestyle="--"),
                    zorder=17)
    ax.text(40, 51.5, "Accès\nchantier", fontsize=6.5, color="#9A7D0A", fontweight="bold",
            ha="center", zorder=33,
            bbox=dict(boxstyle="round,pad=0.1", facecolor="white", alpha=0.85))

    # Barrières + signalisation le long de la route au pont
    for xx in np.linspace(54, 70, 6):
        icon_barrier(ax, xx, 51.2, n=2)
    for xx in [52, 58, 66, 72]:
        ax.plot(xx, 50.3, marker="^", color=RED, ms=8, zorder=26)
        ax.plot(xx, 50.3, marker="^", color="#F7DC6F", ms=4.5, zorder=27)

    # Directions quartiers
    ax.annotate("Vers KIPÉ", xy=(35, 55), xytext=(18, 62),
                fontsize=10, fontweight="bold", color="#9A7D0A",
                arrowprops=dict(arrowstyle="->", color=YELLOW, lw=2.5),
                bbox=dict(boxstyle="round,pad=0.25", facecolor=YELLOW, edgecolor="#9A7D0A"),
                zorder=33)
    ax.annotate("Vers NONGO", xy=(78, 58), xytext=(88, 72),
                fontsize=10, fontweight="bold", color="#9A7D0A",
                arrowprops=dict(arrowstyle="->", color=YELLOW, lw=2.5),
                bbox=dict(boxstyle="round,pad=0.25", facecolor=YELLOW, edgecolor="#9A7D0A"),
                zorder=33)

    # Flèche Nord
    ax.annotate("N", xy=(93, 92), xytext=(93, 84),
                fontsize=14, fontweight="bold", ha="center", color="white",
                arrowprops=dict(arrowstyle="->", color="white", lw=2.5),
                bbox=dict(boxstyle="circle,pad=0.3", facecolor=NAVY, edgecolor="white"),
                zorder=40)

    # Échelle graphique (~ au zoom 18, 5 tiles ≈ zone locale)
    # Approximation : à z18, 1 tile ≈ 150m near equator? Actually at lat 9.6:
    # resolution ≈ 156543 * cos(lat) / 2^z meters/pixel
    # 156543*cos(9.6°)/262144 ≈ 0.59 m/px ; 1280 px ≈ 755 m → 100 units ≈ 755 m → 50 m ≈ 6.6 units
    ax.plot([78, 84.6], [8, 8], color="white", lw=5, zorder=40)
    ax.plot([78, 84.6], [8, 8], color=NAVY, lw=2.5, zorder=41)
    ax.plot([78, 78], [7.3, 8.7], color=NAVY, lw=2, zorder=41)
    ax.plot([84.6, 84.6], [7.3, 8.7], color=NAVY, lw=2, zorder=41)
    ax.text(81.3, 6.2, "0          50 m", fontsize=8, ha="center", color="white",
            fontweight="bold",
            bbox=dict(boxstyle="round,pad=0.15", facecolor=NAVY, alpha=0.85), zorder=42)

    # === BANDEAU TITRE (haut) ===
    ax.add_patch(Rectangle((0, 94.5), 100, 5.5, facecolor=NAVY, edgecolor="none", zorder=50, alpha=0.92))
    ax.text(50, 97.8, "PLAN D'INSTALLATION DU CHANTIER", fontsize=16,
            ha="center", va="center", color="white", fontweight="bold", zorder=51)
    ax.text(50, 95.7,
            "PROJET : CONCEPTION ET DIMENSIONNEMENT D'UN PONT EN BÉTON ARMÉ  —  "
            "PONT KIRIDI ENTRE KIPÉ ET NONGO",
            fontsize=8, ha="center", va="center", color="#F7DC6F", zorder=51)

    # === LÉGENDE (bas droite) ===
    ax.add_patch(FancyBboxPatch((68, 10.5), 30.5, 28, boxstyle="round,pad=0.2,rounding_size=0.5",
                                facecolor="white", edgecolor=NAVY, lw=1.5, alpha=0.93, zorder=45))
    ax.text(83.25, 37, "LÉGENDE", fontsize=10, ha="center", fontweight="bold", color=NAVY, zorder=46)
    legend_items = [
        (1, "#5DADE2", "Base vie"),
        (2, "#F7DC6F", "Magasin de stockage"),
        (3, "#82E0AA", "Aire de ferraillage"),
        (4, "#BB8FCE", "Stockage agrégats"),
        (5, "#5DADE2", "Aire de préfabrication"),
        (6, "#E67E22", "Stationnement camions"),
        (7, "#A04000", "Dépôt des déblais"),
        (8, "#E67E22", "Équipements annexes"),
        (9, "#196F3D", "Point de rassemblement"),
    ]
    for i, (n, c, lab) in enumerate(legend_items):
        yy = 34.5 - i * 2.15
        ax.add_patch(Circle((70.2, yy), 0.7, facecolor=c, edgecolor=NAVY, lw=0.8, zorder=46))
        ax.text(70.2, yy, str(n), ha="center", va="center", fontsize=7, color="white",
                fontweight="bold", zorder=47)
        ax.text(71.5, yy, lab, va="center", fontsize=7.5, color="#1C2833", zorder=46)

    # Symboles complémentaires
    ax.add_patch(FancyBboxPatch((68.5, 11.2), 29, 5.5, boxstyle="round,pad=0.1",
                                facecolor="#F8F9F9", edgecolor="#BFC9CA", lw=0.8, zorder=46))
    ax.plot([70, 74], [14.8, 14.8], color=RED, lw=1.8, ls="--", zorder=47)
    ax.text(74.5, 14.8, "Limite du chantier", va="center", fontsize=6.5, zorder=47)
    ax.plot([70, 74], [13.2, 13.2], color=YELLOW, lw=2.2, zorder=47)
    ax.text(74.5, 13.2, "Déviation / Accès", va="center", fontsize=6.5, zorder=47)
    ax.plot(70.5, 11.8, marker="^", color=RED, ms=7, zorder=47)
    ax.text(74.5, 11.8, "Signalisation / Barrières", va="center", fontsize=6.5, zorder=47)

    # Notes bas gauche
    ax.add_patch(FancyBboxPatch((1.5, 1.5), 48, 7.5, boxstyle="round,pad=0.2",
                                facecolor="white", edgecolor=NAVY, lw=1.2, alpha=0.92, zorder=45))
    ax.text(25.5, 7.2, "NOTES", fontsize=8, ha="center", fontweight="bold", color=NAVY, zorder=46)
    ax.text(25.5, 4.2,
            "• Installations conformes aux règles de sécurité chantier (balisage, HSE).\n"
            "• Implantation adaptée au tissu urbain dense Kipé / Nongo et aux emprises disponibles.\n"
            "• Localisation : ≈ 9°36'41.6″ N – 13°38'11.3″ O  |  Alt. ≈ 3 m  |  Ratoma, Conakry.\n"
            "• Mamady KABA & T.I.H. Diallo  —  Enc. M. Fodé Bangaly Keita.",
            fontsize=6.5, ha="center", va="center", color="#2C3E50", zorder=46)

    fig.tight_layout(pad=0.1)
    path = OUT / "06_PIC_modele_satellite_Kiridi.png"
    fig.savefig(path, dpi=220, bbox_inches="tight", facecolor="white", pad_inches=0.05)
    plt.close(fig)
    return path


def build_deliverables(main_png):
    # PDF single + multi with previous plates if useful
    pdf = Path("/workspace/exports/Plan_Installation_Chantier_Pont_Kiridi.pdf")
    extras = [
        OUT / "04_localisation_Pont_Kiridi.png",
        OUT / "02_organisation_moyens.png",
        OUT / "03_phasage_installation.png",
        OUT / "05_principe_implantation_rives.png",
    ]
    imgs = [PILImage.open(main_png).convert("RGB")]
    for p in extras:
        if p.exists():
            imgs.append(PILImage.open(p).convert("RGB"))
    imgs[0].save(pdf, save_all=True, append_images=imgs[1:], resolution=200)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    NAVY_C = RGBColor(10, 37, 64)
    WHITE = RGBColor(255, 255, 255)

    def slide_img(title, img):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        bar = s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0),
                                 prs.slide_width, Inches(0.65))
        bar.fill.solid(); bar.fill.fore_color.rgb = NAVY_C; bar.line.fill.background()
        tf = s.shapes.add_textbox(Inches(0.3), Inches(0.12), Inches(12.7), Inches(0.4)).text_frame
        p = tf.paragraphs[0]; p.text = title; p.font.size = Pt(15); p.font.bold = True; p.font.color.rgb = WHITE
        s.shapes.add_picture(str(img), Inches(0.5), Inches(0.8), height=Inches(6.5))

    # Cover
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0),
                            prs.slide_width, prs.slide_height)
    bg.fill.solid(); bg.fill.fore_color.rgb = RGBColor(247, 244, 239); bg.line.fill.background()
    bar = s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0),
                             prs.slide_width, Inches(0.7))
    bar.fill.solid(); bar.fill.fore_color.rgb = NAVY_C; bar.line.fill.background()
    tf = s.shapes.add_textbox(Inches(0.4), Inches(0.15), Inches(12.5), Inches(0.45)).text_frame
    p = tf.paragraphs[0]; p.text = "PLAN D'INSTALLATION DU CHANTIER — PONT KIRIDI"
    p.font.size = Pt(18); p.font.bold = True; p.font.color.rgb = WHITE
    body = s.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.5), Inches(4.5)).text_frame
    for i, t in enumerate([
        "Style modèle PIC sur fond satellite réel (Kipé – Nongo, Conakry)",
        "Zones 1 à 9 • Déviation provisoire • Accès • Signalisation • Légende",
        "Coordonnées ≈ 9°36'41.6″ N / 13°38'11.3″ O",
        "",
        "Réalisé par : Mamady KABA & Thierno Ibrahima Hassanatou Diallo",
        "Encadré par : Monsieur Fodé Bangaly Keita",
    ]):
        pp = body.paragraphs[0] if i == 0 else body.add_paragraph()
        pp.text = t
        pp.font.size = Pt(18 if i == 0 else 15)
        pp.font.color.rgb = NAVY_C if i < 3 else RGBColor(80, 80, 80)

    slide_img("PIC — Vue satellite annotée (modèle demandé)", main_png)
    for title, pth in [
        ("Localisation Ratoma / Kipé–Nongo", extras[0]),
        ("Organisation & moyens", extras[1]),
        ("Phasage", extras[2]),
        ("Principe d'implantation des rives", extras[3]),
    ]:
        if pth.exists():
            slide_img(title, pth)

    pptx = Path("/workspace/exports/Plan_Installation_Chantier_Pont_Kiridi.pptx")
    prs.save(pptx)
    return pptx, pdf


def main():
    main_png = fig_pic_modele()
    print(main_png)
    pptx, pdf = build_deliverables(main_png)
    # copy artifacts
    for p in [main_png, pptx, pdf, BASE]:
        if p.exists():
            (ART / p.name).write_bytes(p.read_bytes())
    print(pptx)
    print(pdf)


if __name__ == "__main__":
    main()
