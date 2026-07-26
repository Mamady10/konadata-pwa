#!/usr/bin/env python3
"""Conceptions architecturales - Variantes pont (tablier mixte & béton précontraint)."""

from pathlib import Path
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
from matplotlib.patches import FancyBboxPatch, Rectangle, Polygon, Circle, Arc, FancyArrowPatch
from matplotlib.lines import Line2D
import numpy as np
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from PIL import Image as PILImage

OUT = Path("/workspace/exports/conceptions-variantes")
OUT.mkdir(parents=True, exist_ok=True)

# --- Données projet ---
L_SPAN = 15.70
N_SPANS = 4
L_TOTAL = L_SPAN * N_SPANS  # 62.80 m
NPHE = 4.00
NPBE = 1.00
AFFOUILLEMENT = 5.00
H_CULEE = 5.00
W_CHAUSSEE = 6.00
W_TROTTOIR = 1.00
W_TOTAL = W_CHAUSSEE + 2 * W_TROTTOIR  # 8.00 m

NAVY = "#0A2540"
STEEL = "#4A5568"
CONCRETE = "#C5CED6"
STEEL_FILL = "#6B7C8F"
WATER = "#7EB8DA"
SOIL = "#8B6914"
ACCENT = "#C45C26"
PC_BEAM = "#9AA8B5"
GREEN = "#2F6B4F"


def set_style(ax):
    ax.set_aspect("equal")
    ax.axis("off")
    for spine in ax.spines.values():
        spine.set_visible(False)


def draw_terrain(ax, x0, x1, y_ground_fn, fill=True):
    xs = np.linspace(x0, x1, 200)
    ys = np.array([y_ground_fn(x) for x in xs])
    if fill:
        ax.fill_between(xs, ys, ys.min() - 3, color=SOIL, alpha=0.55, zorder=1)
        ax.fill_between(xs, ys, ys.min() - 3, color="#A67C3D", alpha=0.25, zorder=1)
    ax.plot(xs, ys, color="#5C4033", lw=1.8, zorder=2)


def ground_profile(x, abut_l=0, abut_r=62.8):
    """Profil du terrain sous le pont (vallée)."""
    # Référentiel: x=0 à x=62.8 = axes des culées approx
    mid = (abut_l + abut_r) / 2
    # Dépression centrale
    base = 0.0
    depth = 3.2 * np.exp(-((x - mid) / 14) ** 2)
    # Rives plus hautes
    bank = 0.8 * (np.exp(-((x - abut_l) / 8) ** 2) + np.exp(-((x - abut_r) / 8) ** 2))
    return base - depth + bank * 0.3


def draw_water(ax, x0, x1, y_bed):
    ax.axhline(NPBE, color="#3A8FBF", ls="--", lw=1.0, alpha=0.7, zorder=3)
    ax.axhline(NPHE, color="#1E5F8A", ls="--", lw=1.2, alpha=0.8, zorder=3)
    mid = (x0 + x1) / 2
    wx0, wx1 = mid - 10, mid + 10
    ax.fill_between([wx0, wx1], [y_bed, y_bed], [NPBE, NPBE], color=WATER, alpha=0.45, zorder=2)
    ax.text(-3.5, NPHE + 0.15, f"NPHE = {NPHE:.2f} m", fontsize=8, color="#1E5F8A", ha="left")
    ax.text(-3.5, NPBE + 0.15, f"NPBE = {NPBE:.2f} m", fontsize=8, color="#3A8FBF", ha="left")
    ax.axhline(-AFFOUILLEMENT, color=ACCENT, ls=":", lw=1.2, zorder=3)
    ax.text(-3.5, -AFFOUILLEMENT + 0.25, f"Affouillement prévu = {AFFOUILLEMENT:.2f} m",
            fontsize=8, color=ACCENT, ha="left")


def pier_positions():
    # Culées aux extrémités, piles à 15.7, 31.4, 47.1
    abut_l, abut_r = 0.0, L_TOTAL
    piers = [L_SPAN * i for i in range(1, N_SPANS)]
    return abut_l, abut_r, piers


def draw_dimension_h(ax, x0, x1, y, label, color=NAVY, offset=0.35):
    ax.annotate("", xy=(x1, y), xytext=(x0, y),
                arrowprops=dict(arrowstyle="<->", color=color, lw=1.1))
    ax.text((x0 + x1) / 2, y + offset, label, ha="center", va="bottom",
            fontsize=8, color=color, fontweight="bold")


def draw_dimension_v(ax, x, y0, y1, label, color=NAVY, offset=0.4):
    ax.annotate("", xy=(x, y1), xytext=(x, y0),
                arrowprops=dict(arrowstyle="<->", color=color, lw=1.1))
    ax.text(x + offset, (y0 + y1) / 2, label, ha="left", va="center",
            fontsize=8, color=color, fontweight="bold", rotation=90)


# =============================================================================
# VARIANTE 1 — TABLIER MIXTE
# =============================================================================

def fig_mixte_longitudinal():
    fig, ax = plt.subplots(figsize=(16, 7), dpi=200)
    fig.patch.set_facecolor("#F7F5F1")
    ax.set_facecolor("#F7F5F1")
    set_style(ax)

    abut_l, abut_r, piers = pier_positions()
    deck_y = 6.2
    deck_h = 0.28  # dalle béton
    beam_h = 0.95  # hauteur poutre acier
    slab_top = deck_y + deck_h

    # Terrain
    draw_terrain(ax, -4, L_TOTAL + 4, ground_profile)
    draw_water(ax, 0, L_TOTAL, ground_profile(L_TOTAL / 2) - 0.5)

    # Culées
    for x, side in [(abut_l, "L"), (abut_r, "R")]:
        # Mur de culée
        w = 1.2
        h = H_CULEE + 0.5
        fx = x - w if side == "L" else x
        rect = FancyBboxPatch((fx, deck_y - h), w, h + deck_h,
                              boxstyle="square,pad=0", facecolor=CONCRETE,
                              edgecolor=NAVY, lw=1.4, zorder=5)
        ax.add_patch(rect)
        # Semelle
        sf = FancyBboxPatch((fx - 0.8 if side == "L" else fx - 0.6, deck_y - h - 1.0),
                            w + 1.4, 1.0, boxstyle="square,pad=0",
                            facecolor="#A8B4BE", edgecolor=NAVY, lw=1.2, zorder=4)
        ax.add_patch(sf)
        ax.text(x + (-0.8 if side == "L" else 0.3), deck_y - h / 2,
                f"Culée\nH={H_CULEE:.0f}m", fontsize=7.5, color=NAVY, ha="center",
                fontweight="bold", zorder=6)

    # Piles + chevêtres
    for i, px in enumerate(piers):
        g = ground_profile(px)
        # Semelle plus profonde au centre
        depth = 2.2 if i == 1 else 1.6
        foot_y = g - depth
        foot = FancyBboxPatch((px - 1.4, foot_y), 2.8, 0.9,
                              boxstyle="square,pad=0", facecolor="#A8B4BE",
                              edgecolor=NAVY, lw=1.2, zorder=4)
        ax.add_patch(foot)
        # Fût de pile
        pier_top = deck_y - beam_h - 0.35
        pier = FancyBboxPatch((px - 0.55, foot_y + 0.9), 1.1, pier_top - (foot_y + 0.9),
                              boxstyle="square,pad=0", facecolor=CONCRETE,
                              edgecolor=NAVY, lw=1.3, zorder=5)
        ax.add_patch(pier)
        # Chevêtre
        chev = FancyBboxPatch((px - 1.5, pier_top - 0.05), 3.0, 0.55,
                              boxstyle="square,pad=0", facecolor="#B8C4CE",
                              edgecolor=NAVY, lw=1.2, zorder=6)
        ax.add_patch(chev)
        # Appareils d'appui
        for dx in [-0.7, 0.7]:
            ax.add_patch(Rectangle((px + dx - 0.18, pier_top + 0.5), 0.36, 0.12,
                                   facecolor="#333", edgecolor="none", zorder=7))
        ax.text(px, foot_y - 0.35, "Semelle", fontsize=6.5, ha="center", color=NAVY)

    # Poutres acier (profil en I schématisé) + dalle
    # Ãme acier sous la dalle sur toute la longueur
    ax.add_patch(Rectangle((abut_l, deck_y - beam_h), L_TOTAL, beam_h,
                           facecolor=STEEL_FILL, edgecolor=STEEL, lw=1.3, zorder=8))
    # Semelles acier haut/bas (flanges)
    ax.add_patch(Rectangle((abut_l, deck_y - 0.12), L_TOTAL, 0.12,
                           facecolor=STEEL, edgecolor=STEEL, lw=0.5, zorder=9))
    ax.add_patch(Rectangle((abut_l, deck_y - beam_h), L_TOTAL, 0.12,
                           facecolor=STEEL, edgecolor=STEEL, lw=0.5, zorder=9))

    # Dalle béton
    ax.add_patch(Rectangle((abut_l - 0.3, deck_y), L_TOTAL + 0.6, deck_h,
                           facecolor="#D6DDE4", edgecolor=NAVY, lw=1.4, zorder=10))
    # Garde-corps
    for x in np.linspace(0, L_TOTAL, 25):
        ax.plot([x, x], [slab_top, slab_top + 0.95], color=NAVY, lw=0.8, zorder=11)
    ax.plot([0, L_TOTAL], [slab_top + 0.95, slab_top + 0.95], color=NAVY, lw=1.2, zorder=11)

    # Dalles de transition
    ax.add_patch(Rectangle((-3.0, deck_y), 2.7, deck_h * 0.85,
                           facecolor="#C8D0D8", edgecolor=NAVY, lw=1.0, zorder=9))
    ax.add_patch(Rectangle((L_TOTAL + 0.3, deck_y), 2.7, deck_h * 0.85,
                           facecolor="#C8D0D8", edgecolor=NAVY, lw=1.0, zorder=9))

    # Cotes portées
    for i in range(N_SPANS):
        x0 = i * L_SPAN
        x1 = (i + 1) * L_SPAN
        draw_dimension_h(ax, x0, x1, slab_top + 1.6, f"{L_SPAN:.2f} m")
    draw_dimension_h(ax, 0, L_TOTAL, slab_top + 2.35, f"L totale = {L_TOTAL:.2f} m", color=ACCENT)

    # Légende éléments
    ax.text(L_TOTAL / 2, deck_y + 0.08, "Dalle BA  (hourdis)", fontsize=8,
            ha="center", color=NAVY, fontweight="bold", zorder=12)
    ax.text(L_TOTAL / 2, deck_y - beam_h / 2, "Poutres en acier  (profilés I)",
            fontsize=8, ha="center", va="center", color="white", fontweight="bold", zorder=12)

    ax.set_xlim(-6, L_TOTAL + 6)
    ax.set_ylim(-6.5, slab_top + 3.2)
    ax.set_title("VARIANTE A — Pont à tablier mixte (béton armé + poutres acier)\n"
                 "Profil longitudinal  |  Portée 4 × 15,70 m = 62,80 m",
                 fontsize=13, fontweight="bold", color=NAVY, pad=12)

    # Cartouche
    ax.text(0.01, 0.02,
            "Mémoire d'ingénieurs | Ponts et chaussées | Mamady KABA & T.I.H. Diallo  |  Enc. F.B. Keita",
            transform=ax.transAxes, fontsize=7.5, color="#555")
    fig.tight_layout()
    path = OUT / "01_mixte_profil_longitudinal.png"
    fig.savefig(path, dpi=220, bbox_inches="tight", facecolor=fig.get_facecolor())
    plt.close(fig)
    return path


def fig_mixte_plan():
    fig, ax = plt.subplots(figsize=(16, 5.5), dpi=200)
    fig.patch.set_facecolor("#F7F5F1")
    ax.set_facecolor("#F7F5F1")
    set_style(ax)

    W = W_TOTAL
    # Tablier
    ax.add_patch(FancyBboxPatch((0, 0), L_TOTAL, W, boxstyle="square,pad=0",
                                facecolor="#D9E0E7", edgecolor=NAVY, lw=1.6, zorder=3))
    # Chaussée
    ax.add_patch(Rectangle((0, W_TROTTOIR), L_TOTAL, W_CHAUSSEE,
                           facecolor="#6E7680", edgecolor="none", zorder=4, alpha=0.55))
    # Ligne médiane
    for x in np.arange(1, L_TOTAL, 2.0):
        ax.plot([x, x + 1.0], [W / 2, W / 2], color="white", lw=1.5, zorder=5, ls="--")

    # Poutres acier (vues en plan = 5 files)
    n_beams = 5
    spacing = W_CHAUSSEE / (n_beams - 1)
    y0 = W_TROTTOIR
    for i in range(n_beams):
        y = y0 + i * spacing
        ax.plot([0.3, L_TOTAL - 0.3], [y, y], color=STEEL, lw=2.4, zorder=6)
        if i == 0:
            ax.text(L_TOTAL / 2, y - 0.35, "Axe poutre acier", fontsize=7,
                    ha="center", color=STEEL)

    # Piles + culées
    abut_l, abut_r, piers = pier_positions()
    for px in [abut_l, abut_r]:
        ax.add_patch(Rectangle((px - 0.8, -1.2), 1.6, W + 2.4,
                               facecolor=CONCRETE, edgecolor=NAVY, lw=1.2, zorder=2, alpha=0.9))
    for px in piers:
        ax.add_patch(Circle((px, W / 2), 0.7, facecolor=CONCRETE, edgecolor=NAVY, lw=1.3, zorder=7))
        ax.add_patch(FancyBboxPatch((px - 1.6, -0.6), 3.2, W + 1.2, boxstyle="square,pad=0",
                                    fill=False, edgecolor=NAVY, lw=1.0, ls="--", zorder=2))
        ax.text(px, W + 1.0, "Pile", fontsize=8, ha="center", color=NAVY, fontweight="bold")

    ax.text(-0.2, W / 2, "Culée C0", fontsize=8, ha="right", va="center", color=NAVY, fontweight="bold")
    ax.text(L_TOTAL + 0.2, W / 2, "Culée C1", fontsize=8, ha="left", va="center", color=NAVY, fontweight="bold")

    # Trottoirs labels
    ax.text(L_TOTAL / 2, W_TROTTOIR / 2, "Trottoir 1,00 m", fontsize=8, ha="center", color=NAVY)
    ax.text(L_TOTAL / 2, W - W_TROTTOIR / 2, "Trottoir 1,00 m", fontsize=8, ha="center", color=NAVY)
    ax.text(L_TOTAL / 2, W / 2 + 0.9, "Chaussée 6,00 m (2 voies)", fontsize=9,
            ha="center", color="white", fontweight="bold", zorder=8)

    # Cotes
    draw_dimension_h(ax, 0, L_TOTAL, W + 2.0, f"{L_TOTAL:.2f} m")
    for i in range(N_SPANS):
        draw_dimension_h(ax, i * L_SPAN, (i + 1) * L_SPAN, -1.8, f"{L_SPAN:.2f} m")
    draw_dimension_v(ax, L_TOTAL + 3.2, 0, W, f"{W:.2f} m")

    # Murs de soutènement
    ax.add_patch(Rectangle((-2.5, -1.0), 2.5, 1.0, facecolor="#B0B8C0", edgecolor=NAVY, lw=1.0))
    ax.add_patch(Rectangle((-2.5, W), 2.5, 1.0, facecolor="#B0B8C0", edgecolor=NAVY, lw=1.0))
    ax.text(-1.2, -0.5, "Mur de\nsoutènement", fontsize=6.5, ha="center", va="center", color=NAVY)

    ax.set_xlim(-5, L_TOTAL + 5)
    ax.set_ylim(-3.2, W + 3.0)
    ax.set_title("VARIANTE A — Vue en plan  |  Tablier mixte  |  Largeur utile 8,00 m",
                 fontsize=13, fontweight="bold", color=NAVY, pad=10)
    fig.tight_layout()
    path = OUT / "02_mixte_vue_en_plan.png"
    fig.savefig(path, dpi=220, bbox_inches="tight", facecolor=fig.get_facecolor())
    plt.close(fig)
    return path


def fig_mixte_coupe():
    fig, ax = plt.subplots(figsize=(12, 8), dpi=200)
    fig.patch.set_facecolor("#F7F5F1")
    ax.set_facecolor("#F7F5F1")
    set_style(ax)

    # Coupe transversale du tablier mixte
    # Origine au centre de la chaussée
    half = W_TOTAL / 2
    slab_t = 0.22
    haunch = 0.08
    beam_h = 0.90
    flange_w = 0.35
    flange_t = 0.04
    web_t = 0.016

    # Dalle
    ax.add_patch(Rectangle((-half - 0.15, 0), W_TOTAL + 0.3, slab_t,
                           facecolor="#D6DDE4", edgecolor=NAVY, lw=1.5, zorder=5))
    # Corniche / trottoir surélévation légère
    for sgn in [-1, 1]:
        x0 = sgn * (half - W_TROTTOIR)
        ax.add_patch(Rectangle((min(x0, sgn * half) - (0.05 if sgn < 0 else 0),
                                slab_t),
                               W_TROTTOIR + 0.05, 0.12,
                               facecolor="#C5CED6", edgecolor=NAVY, lw=1.0, zorder=6))

    # Chaussée marquage
    ax.plot([-W_CHAUSSEE / 2, W_CHAUSSEE / 2], [slab_t + 0.01, slab_t + 0.01],
            color="#6E7680", lw=8, solid_capstyle="butt", zorder=4, alpha=0.35)

    # 5 poutres I acier
    n = 5
    positions = np.linspace(-W_CHAUSSEE / 2, W_CHAUSSEE / 2, n)
    for i, xc in enumerate(positions):
        # Flange top
        ax.add_patch(Rectangle((xc - flange_w / 2, -flange_t), flange_w, flange_t,
                               facecolor=STEEL, edgecolor="#222", lw=0.8, zorder=7))
        # Web
        ax.add_patch(Rectangle((xc - web_t / 2, -beam_h + flange_t), web_t, beam_h - 2 * flange_t,
                               facecolor=STEEL_FILL, edgecolor="#222", lw=0.6, zorder=7))
        # Flange bottom
        ax.add_patch(Rectangle((xc - flange_w / 2, -beam_h), flange_w, flange_t,
                               facecolor=STEEL, edgecolor="#222", lw=0.8, zorder=7))
        # Connecteurs (goujons)
        for dx in [-0.08, 0.08]:
            ax.plot([xc + dx, xc + dx], [0, slab_t * 0.7], color=ACCENT, lw=1.5, zorder=8)
            ax.add_patch(Circle((xc + dx, slab_t * 0.65), 0.025, facecolor=ACCENT, zorder=9))
        ax.text(xc, -beam_h - 0.25, f"P{i+1}", fontsize=8, ha="center", color=STEEL, fontweight="bold")

    # Garde-corps
    for sgn in [-1, 1]:
        x = sgn * half
        ax.plot([x, x], [slab_t + 0.12, slab_t + 1.15], color=NAVY, lw=1.6, zorder=10)
        ax.plot([x - 0.12 * sgn, x], [slab_t + 1.15, slab_t + 1.15], color=NAVY, lw=1.4, zorder=10)

    # Entretoise
    ax.plot([positions[0], positions[-1]], [-beam_h / 2, -beam_h / 2],
            color="#405060", lw=1.8, ls="-", zorder=6)
    ax.text(0, -beam_h / 2 + 0.12, "Entretoise métallique", fontsize=8,
            ha="center", color="#405060")

    # Annotations
    draw_dimension_h(ax, -half, half, slab_t + 1.55, f"Largeur totale = {W_TOTAL:.2f} m")
    draw_dimension_h(ax, -W_CHAUSSEE / 2, W_CHAUSSEE / 2, -beam_h - 0.7,
                     f"Chaussée = {W_CHAUSSEE:.2f} m")
    draw_dimension_v(ax, half + 0.9, -beam_h, 0, f"h poutre ≈ {beam_h:.2f} m")
    draw_dimension_v(ax, half + 1.7, 0, slab_t, f"e dalle = {slab_t*100:.0f} cm")

    # Cartouche légende
    ax.add_patch(FancyBboxPatch((-half - 0.3, -beam_h - 2.0), W_TOTAL + 0.6, 1.0,
                                boxstyle="round,pad=0.05", facecolor="white",
                                edgecolor=NAVY, lw=1.0, zorder=3))
    ax.text(0, -beam_h - 1.35,
            "Goujons de connexion  •  Poutre I acier  •  Dalle BA collaborante  •  Action mixte",
            fontsize=9, ha="center", color=NAVY, fontweight="bold")
    ax.text(0, -beam_h - 1.75,
            "Avantages : allègement du tablier, grandes portées, montage rapide, faible hauteur constructive",
            fontsize=8, ha="center", color="#444")

    ax.set_xlim(-half - 2.5, half + 3.2)
    ax.set_ylim(-beam_h - 2.4, slab_t + 2.0)
    ax.set_title("VARIANTE A — Coupe transversale du tablier mixte\n"
                 "(dalle en béton armé collaborante + 5 poutres en acier)",
                 fontsize=12, fontweight="bold", color=NAVY, pad=10)
    fig.tight_layout()
    path = OUT / "03_mixte_coupe_transversale.png"
    fig.savefig(path, dpi=220, bbox_inches="tight", facecolor=fig.get_facecolor())
    plt.close(fig)
    return path


def fig_mixte_detail():
    fig, ax = plt.subplots(figsize=(11, 7), dpi=200)
    fig.patch.set_facecolor("#F7F5F1")
    ax.set_facecolor("#F7F5F1")
    set_style(ax)

    # Détail connexion mixte à l'échelle agrandie
    # Dalle
    ax.add_patch(Rectangle((-1.2, 0), 2.4, 0.5, facecolor="#D6DDE4", edgecolor=NAVY, lw=1.6))
    # Armatures dalle
    for y in [0.15, 0.35]:
        ax.plot([-1.1, 1.1], [y, y], color="#888", lw=1.2, ls="-")
    ax.text(0, 0.55, "Dalle en béton armé (collaborante)", fontsize=10,
            ha="center", color=NAVY, fontweight="bold")

    # Semelle supérieure acier
    ax.add_patch(Rectangle((-0.55, -0.12), 1.1, 0.12, facecolor=STEEL, edgecolor="#111", lw=1.2))
    # Âme
    ax.add_patch(Rectangle((-0.04, -1.4), 0.08, 1.28, facecolor=STEEL_FILL, edgecolor="#111", lw=1.0))
    # Semelle inférieure
    ax.add_patch(Rectangle((-0.55, -1.52), 1.1, 0.12, facecolor=STEEL, edgecolor="#111", lw=1.2))

    # Goujons Nelson
    for x in [-0.25, 0.0, 0.25]:
        ax.plot([x, x], [0, 0.35], color=ACCENT, lw=3.0, zorder=5)
        ax.add_patch(Circle((x, 0.38), 0.05, facecolor=ACCENT, edgecolor="#8B3010", zorder=6))
    ax.annotate("Goujons de cisaillement\n(connexion mixte)",
                xy=(0.25, 0.25), xytext=(1.0, 1.0),
                fontsize=9, color=ACCENT, fontweight="bold",
                arrowprops=dict(arrowstyle="->", color=ACCENT, lw=1.3))

    ax.annotate("Semelle supérieure", xy=(0.55, -0.06), xytext=(1.2, -0.4),
                fontsize=9, color=STEEL,
                arrowprops=dict(arrowstyle="->", color=STEEL))
    ax.annotate("Âme de la poutre", xy=(0.04, -0.7), xytext=(1.2, -0.9),
                fontsize=9, color=STEEL,
                arrowprops=dict(arrowstyle="->", color=STEEL))
    ax.annotate("Semelle inférieure", xy=(0.55, -1.46), xytext=(1.2, -1.7),
                fontsize=9, color=STEEL,
                arrowprops=dict(arrowstyle="->", color=STEEL))

    # Chevêtre + appui
    ax.add_patch(FancyBboxPatch((-1.6, -2.4), 3.2, 0.55, boxstyle="square,pad=0",
                                facecolor=CONCRETE, edgecolor=NAVY, lw=1.3))
    ax.add_patch(Rectangle((-0.35, -1.85), 0.7, 0.25, facecolor="#333", edgecolor="#111", lw=1.0))
    ax.text(0, -2.05, "Appareil d'appui\n(élastomère)", fontsize=8, ha="center", color="white",
            fontweight="bold", va="center")
    ax.text(0, -2.15 - 0.35, "Chevêtre BA", fontsize=9, ha="center", color=NAVY, fontweight="bold")

    ax.set_xlim(-2.2, 3.0)
    ax.set_ylim(-2.9, 1.5)
    ax.set_title("VARIANTE A — Détail de connexion mixte acier/béton\n"
                 "et appui sur chevêtre",
                 fontsize=12, fontweight="bold", color=NAVY, pad=10)
    fig.tight_layout()
    path = OUT / "04_mixte_detail_connexion.png"
    fig.savefig(path, dpi=220, bbox_inches="tight", facecolor=fig.get_facecolor())
    plt.close(fig)
    return path


# =============================================================================
# VARIANTE 2 — BÉTON PRÉCONTRAINT
# =============================================================================

def fig_pc_longitudinal():
    fig, ax = plt.subplots(figsize=(16, 7), dpi=200)
    fig.patch.set_facecolor("#F4F7F5")
    ax.set_facecolor("#F4F7F5")
    set_style(ax)

    abut_l, abut_r, piers = pier_positions()
    deck_y = 6.2
    beam_h = 1.15  # poutres précontraintes plus hautes
    slab_t = 0.20
    slab_top = deck_y + slab_t

    draw_terrain(ax, -4, L_TOTAL + 4, ground_profile)
    draw_water(ax, 0, L_TOTAL, ground_profile(L_TOTAL / 2) - 0.5)

    # Culées
    for x, side in [(abut_l, "L"), (abut_r, "R")]:
        w = 1.25
        h = H_CULEE + 0.5
        fx = x - w if side == "L" else x
        ax.add_patch(FancyBboxPatch((fx, deck_y - h), w, h + slab_t,
                                    boxstyle="square,pad=0", facecolor=CONCRETE,
                                    edgecolor=GREEN, lw=1.5, zorder=5))
        ax.add_patch(FancyBboxPatch((fx - 0.8 if side == "L" else fx - 0.6,
                                     deck_y - h - 1.0), w + 1.4, 1.0,
                                    boxstyle="square,pad=0", facecolor="#A8B4BE",
                                    edgecolor=GREEN, lw=1.2, zorder=4))
        ax.text(x + (-0.85 if side == "L" else 0.35), deck_y - h / 2,
                f"Culée\nH={H_CULEE:.0f}m", fontsize=7.5, color=GREEN,
                ha="center", fontweight="bold", zorder=6)

    for i, px in enumerate(piers):
        g = ground_profile(px)
        depth = 2.2 if i == 1 else 1.6
        foot_y = g - depth
        ax.add_patch(FancyBboxPatch((px - 1.4, foot_y), 2.8, 0.9,
                                    boxstyle="square,pad=0", facecolor="#A8B4BE",
                                    edgecolor=GREEN, lw=1.2, zorder=4))
        pier_top = deck_y - beam_h - 0.35
        ax.add_patch(FancyBboxPatch((px - 0.55, foot_y + 0.9), 1.1,
                                    pier_top - (foot_y + 0.9),
                                    boxstyle="square,pad=0", facecolor=CONCRETE,
                                    edgecolor=GREEN, lw=1.3, zorder=5))
        ax.add_patch(FancyBboxPatch((px - 1.55, pier_top - 0.05), 3.1, 0.55,
                                    boxstyle="square,pad=0", facecolor="#B8C4CE",
                                    edgecolor=GREEN, lw=1.2, zorder=6))
        for dx in [-0.7, 0.7]:
            ax.add_patch(Rectangle((px + dx - 0.18, pier_top + 0.5), 0.36, 0.12,
                                   facecolor="#333", zorder=7))

    # Poutres précontraintes (par travée, discontinuïté visuelle aux appuis)
    for i in range(N_SPANS):
        x0 = i * L_SPAN + 0.15
        x1 = (i + 1) * L_SPAN - 0.15
        # Corps poutre
        ax.add_patch(FancyBboxPatch((x0, deck_y - beam_h), x1 - x0, beam_h,
                                    boxstyle="square,pad=0", facecolor=PC_BEAM,
                                    edgecolor=GREEN, lw=1.4, zorder=8))
        # Câbles de précontrainte (parabole schématisée)
        xs = np.linspace(x0 + 0.3, x1 - 0.3, 40)
        mid = (x0 + x1) / 2
        ys = deck_y - 0.25 - 0.65 * (1 - ((xs - mid) / ((x1 - x0) / 2)) ** 2)
        ax.plot(xs, ys, color=ACCENT, lw=1.8, zorder=9)
        if i == 1:
            ax.text(mid, deck_y - beam_h / 2 - 0.15, "Câbles de\nprécontrainte",
                    fontsize=7.5, ha="center", color=ACCENT, fontweight="bold", zorder=10)

    # Hourdis coulé en place
    ax.add_patch(Rectangle((abut_l - 0.3, deck_y), L_TOTAL + 0.6, slab_t,
                           facecolor="#D6DDE4", edgecolor=GREEN, lw=1.5, zorder=10))
    for x in np.linspace(0, L_TOTAL, 25):
        ax.plot([x, x], [slab_top, slab_top + 0.95], color=GREEN, lw=0.8, zorder=11)
    ax.plot([0, L_TOTAL], [slab_top + 0.95, slab_top + 0.95], color=GREEN, lw=1.2, zorder=11)

    ax.add_patch(Rectangle((-3.0, deck_y), 2.7, slab_t * 0.85,
                           facecolor="#C8D0D8", edgecolor=GREEN, lw=1.0, zorder=9))
    ax.add_patch(Rectangle((L_TOTAL + 0.3, deck_y), 2.7, slab_t * 0.85,
                           facecolor="#C8D0D8", edgecolor=GREEN, lw=1.0, zorder=9))

    for i in range(N_SPANS):
        draw_dimension_h(ax, i * L_SPAN, (i + 1) * L_SPAN, slab_top + 1.6, f"{L_SPAN:.2f} m", color=GREEN)
    draw_dimension_h(ax, 0, L_TOTAL, slab_top + 2.35, f"L totale = {L_TOTAL:.2f} m", color=ACCENT)

    ax.text(L_TOTAL / 2, deck_y + 0.06, "Hourdis BA coulé en place", fontsize=8,
            ha="center", color=NAVY, fontweight="bold", zorder=12)
    ax.text(L_TOTAL * 0.12, deck_y - beam_h / 2, "Poutre préfabriquée\nen béton précontraint",
            fontsize=7.5, ha="center", va="center", color=GREEN, fontweight="bold", zorder=12)

    ax.set_xlim(-6, L_TOTAL + 6)
    ax.set_ylim(-6.5, slab_top + 3.2)
    ax.set_title("VARIANTE B — Pont en béton précontraint (poutres préfabriquées + hourdis)\n"
                 "Profil longitudinal  |  Portée 4 × 15,70 m = 62,80 m",
                 fontsize=13, fontweight="bold", color=GREEN, pad=12)
    ax.text(0.01, 0.02,
            "Mémoire d'ingénieurs | Ponts et chaussées | Mamady KABA & T.I.H. Diallo  |  Enc. F.B. Keita",
            transform=ax.transAxes, fontsize=7.5, color="#555")
    fig.tight_layout()
    path = OUT / "05_precontraint_profil_longitudinal.png"
    fig.savefig(path, dpi=220, bbox_inches="tight", facecolor=fig.get_facecolor())
    plt.close(fig)
    return path


def fig_pc_plan():
    fig, ax = plt.subplots(figsize=(16, 5.5), dpi=200)
    fig.patch.set_facecolor("#F4F7F5")
    ax.set_facecolor("#F4F7F5")
    set_style(ax)

    W = W_TOTAL
    ax.add_patch(FancyBboxPatch((0, 0), L_TOTAL, W, boxstyle="square,pad=0",
                                facecolor="#D5E0D8", edgecolor=GREEN, lw=1.6, zorder=3))
    ax.add_patch(Rectangle((0, W_TROTTOIR), L_TOTAL, W_CHAUSSEE,
                           facecolor="#6E7680", edgecolor="none", zorder=4, alpha=0.5))
    for x in np.arange(1, L_TOTAL, 2.0):
        ax.plot([x, x + 1.0], [W / 2, W / 2], color="white", lw=1.5, zorder=5, ls="--")

    n_beams = 5
    spacing = W_CHAUSSEE / (n_beams - 1)
    y0 = W_TROTTOIR
    for i in range(n_beams):
        y = y0 + i * spacing
        # Poutres préfabriquées - bande
        ax.add_patch(Rectangle((0.4, y - 0.28), L_TOTAL - 0.8, 0.56,
                               facecolor=PC_BEAM, edgecolor=GREEN, lw=1.1, zorder=6, alpha=0.9))
        if i == 2:
            ax.text(L_TOTAL / 2, y, "Poutre préfabriquée précontrainte",
                    fontsize=7.5, ha="center", va="center", color=GREEN, fontweight="bold", zorder=7)

    abut_l, abut_r, piers = pier_positions()
    for px in [abut_l, abut_r]:
        ax.add_patch(Rectangle((px - 0.8, -1.2), 1.6, W + 2.4,
                               facecolor=CONCRETE, edgecolor=GREEN, lw=1.2, zorder=2))
    for px in piers:
        ax.add_patch(Circle((px, W / 2), 0.7, facecolor=CONCRETE, edgecolor=GREEN, lw=1.3, zorder=7))
        ax.add_patch(FancyBboxPatch((px - 1.6, -0.6), 3.2, W + 1.2, boxstyle="square,pad=0",
                                    fill=False, edgecolor=GREEN, lw=1.0, ls="--", zorder=2))
        ax.text(px, W + 1.0, "Pile", fontsize=8, ha="center", color=GREEN, fontweight="bold")

    ax.text(L_TOTAL / 2, W_TROTTOIR / 2, "Trottoir 1,00 m", fontsize=8, ha="center", color=GREEN)
    ax.text(L_TOTAL / 2, W - W_TROTTOIR / 2, "Trottoir 1,00 m", fontsize=8, ha="center", color=GREEN)
    ax.text(L_TOTAL / 2, W / 2 + 1.1, "Chaussée 6,00 m (2 voies)", fontsize=9,
            ha="center", color="white", fontweight="bold", zorder=8)

    draw_dimension_h(ax, 0, L_TOTAL, W + 2.0, f"{L_TOTAL:.2f} m", color=GREEN)
    for i in range(N_SPANS):
        draw_dimension_h(ax, i * L_SPAN, (i + 1) * L_SPAN, -1.8, f"{L_SPAN:.2f} m", color=GREEN)
    draw_dimension_v(ax, L_TOTAL + 3.2, 0, W, f"{W:.2f} m", color=GREEN)

    ax.set_xlim(-5, L_TOTAL + 5)
    ax.set_ylim(-3.2, W + 3.0)
    ax.set_title("VARIANTE B — Vue en plan  |  Béton précontraint  |  Largeur utile 8,00 m",
                 fontsize=13, fontweight="bold", color=GREEN, pad=10)
    fig.tight_layout()
    path = OUT / "06_precontraint_vue_en_plan.png"
    fig.savefig(path, dpi=220, bbox_inches="tight", facecolor=fig.get_facecolor())
    plt.close(fig)
    return path


def fig_pc_coupe():
    fig, ax = plt.subplots(figsize=(12, 8), dpi=200)
    fig.patch.set_facecolor("#F4F7F5")
    ax.set_facecolor("#F4F7F5")
    set_style(ax)

    half = W_TOTAL / 2
    slab_t = 0.20
    beam_h = 1.10
    # Poutres I précontraintes
    n = 5
    positions = np.linspace(-W_CHAUSSEE / 2, W_CHAUSSEE / 2, n)

    ax.add_patch(Rectangle((-half - 0.15, 0), W_TOTAL + 0.3, slab_t,
                           facecolor="#D6DDE4", edgecolor=GREEN, lw=1.5, zorder=5))
    for sgn in [-1, 1]:
        x0 = sgn * (half - W_TROTTOIR)
        ax.add_patch(Rectangle((min(x0, sgn * half) - (0.05 if sgn < 0 else 0), slab_t),
                               W_TROTTOIR + 0.05, 0.12,
                               facecolor="#C5CED6", edgecolor=GREEN, lw=1.0, zorder=6))

    def draw_i_beam(xc):
        # Forme I typique préfabriquée
        top_w, bot_w = 0.55, 0.70
        top_t, bot_t = 0.14, 0.18
        web_w = 0.16
        # Semelle haute
        ax.add_patch(Polygon([
            (xc - top_w / 2, 0), (xc + top_w / 2, 0),
            (xc + top_w / 2, -top_t), (xc + web_w / 2, -top_t),
            (xc + web_w / 2, -beam_h + bot_t), (xc + bot_w / 2, -beam_h + bot_t),
            (xc + bot_w / 2, -beam_h), (xc - bot_w / 2, -beam_h),
            (xc - bot_w / 2, -beam_h + bot_t), (xc - web_w / 2, -beam_h + bot_t),
            (xc - web_w / 2, -top_t), (xc - top_w / 2, -top_t),
        ], closed=True, facecolor=PC_BEAM, edgecolor=GREEN, lw=1.3, zorder=7))
        # Câbles (cercles dans âme basse)
        for dx, dy in [(-0.05, -beam_h + 0.35), (0.05, -beam_h + 0.35),
                       (0.0, -beam_h + 0.50)]:
            ax.add_patch(Circle((xc + dx, dy), 0.035, facecolor=ACCENT,
                                edgecolor="#8B3010", lw=0.6, zorder=8))

    for i, xc in enumerate(positions):
        draw_i_beam(xc)
        ax.text(xc, -beam_h - 0.28, f"PP{i+1}", fontsize=8, ha="center",
                color=GREEN, fontweight="bold")

    for sgn in [-1, 1]:
        x = sgn * half
        ax.plot([x, x], [slab_t + 0.12, slab_t + 1.15], color=GREEN, lw=1.6, zorder=10)
        ax.plot([x - 0.12 * sgn, x], [slab_t + 1.15, slab_t + 1.15], color=GREEN, lw=1.4, zorder=10)

    # Entretoise béton
    ax.add_patch(Rectangle((positions[0] - 0.1, -beam_h / 2 - 0.12),
                           positions[-1] - positions[0] + 0.2, 0.24,
                           facecolor=CONCRETE, edgecolor=GREEN, lw=1.0, zorder=6, alpha=0.85))
    ax.text(0, -beam_h / 2, "Entretoise BA", fontsize=8, ha="center", va="center",
            color=GREEN, fontweight="bold", zorder=9)

    draw_dimension_h(ax, -half, half, slab_t + 1.55, f"Largeur totale = {W_TOTAL:.2f} m", color=GREEN)
    draw_dimension_h(ax, -W_CHAUSSEE / 2, W_CHAUSSEE / 2, -beam_h - 0.75,
                     f"Chaussée = {W_CHAUSSEE:.2f} m", color=GREEN)
    draw_dimension_v(ax, half + 1.0, -beam_h, 0, f"h poutre ≈ {beam_h:.2f} m", color=GREEN)
    draw_dimension_v(ax, half + 1.8, 0, slab_t, f"e hourdis = {slab_t*100:.0f} cm", color=GREEN)

    ax.add_patch(FancyBboxPatch((-half - 0.3, -beam_h - 2.15), W_TOTAL + 0.6, 1.1,
                                boxstyle="round,pad=0.05", facecolor="white",
                                edgecolor=GREEN, lw=1.0, zorder=3))
    ax.text(0, -beam_h - 1.4,
            "Poutres I préfabriquées précontraintes  •  Câbles (cercles)  •  Hourdis collaborant",
            fontsize=9, ha="center", color=GREEN, fontweight="bold")
    ax.text(0, -beam_h - 1.85,
            "Avantages : contrôle qualité en usine, montage rapide, excellente durabilité, faible entretien",
            fontsize=8, ha="center", color="#444")

    ax.set_xlim(-half - 2.5, half + 3.2)
    ax.set_ylim(-beam_h - 2.5, slab_t + 2.0)
    ax.set_title("VARIANTE B — Coupe transversale\n"
                 "(5 poutres en béton précontraint + hourdis BA)",
                 fontsize=12, fontweight="bold", color=GREEN, pad=10)
    fig.tight_layout()
    path = OUT / "07_precontraint_coupe_transversale.png"
    fig.savefig(path, dpi=220, bbox_inches="tight", facecolor=fig.get_facecolor())
    plt.close(fig)
    return path


def fig_pc_detail():
    fig, ax = plt.subplots(figsize=(11, 7), dpi=200)
    fig.patch.set_facecolor("#F4F7F5")
    ax.set_facecolor("#F4F7F5")
    set_style(ax)

    # Détail poutre I + hourdis + câbles
    # Hourdis
    ax.add_patch(Rectangle((-1.4, 0), 2.8, 0.45, facecolor="#D6DDE4", edgecolor=GREEN, lw=1.6))
    ax.text(0, 0.55, "Hourdis BA coulé en place", fontsize=10, ha="center",
            color=NAVY, fontweight="bold")

    # Poutre I
    top_w, bot_w = 1.0, 1.25
    top_t, bot_t = 0.22, 0.28
    web_w = 0.28
    beam_h = 1.6
    xc = 0
    ax.add_patch(Polygon([
        (xc - top_w / 2, 0), (xc + top_w / 2, 0),
        (xc + top_w / 2, -top_t), (xc + web_w / 2, -top_t),
        (xc + web_w / 2, -beam_h + bot_t), (xc + bot_w / 2, -beam_h + bot_t),
        (xc + bot_w / 2, -beam_h), (xc - bot_w / 2, -beam_h),
        (xc - bot_w / 2, -beam_h + bot_t), (xc - web_w / 2, -beam_h + bot_t),
        (xc - web_w / 2, -top_t), (xc - top_w / 2, -top_t),
    ], closed=True, facecolor=PC_BEAM, edgecolor=GREEN, lw=1.6, zorder=5))

    # Câbles
    cables = [(-0.12, -beam_h + 0.45), (0.12, -beam_h + 0.45),
              (0.0, -beam_h + 0.62), (-0.12, -beam_h + 0.78), (0.12, -beam_h + 0.78)]
    for cx, cy in cables:
        ax.add_patch(Circle((cx, cy), 0.055, facecolor=ACCENT, edgecolor="#8B3010", lw=0.8, zorder=6))

    ax.annotate("Câbles de précontrainte\n(torons)",
                xy=(0.12, -beam_h + 0.45), xytext=(1.4, -0.5),
                fontsize=9, color=ACCENT, fontweight="bold",
                arrowprops=dict(arrowstyle="->", color=ACCENT, lw=1.3))
    ax.annotate("Semelle inférieure", xy=(bot_w / 2, -beam_h + 0.05),
                xytext=(1.5, -beam_h - 0.15),
                fontsize=9, color=GREEN,
                arrowprops=dict(arrowstyle="->", color=GREEN))
    ax.annotate("Âme", xy=(web_w / 2, -0.9), xytext=(1.4, -1.0),
                fontsize=9, color=GREEN,
                arrowprops=dict(arrowstyle="->", color=GREEN))

    # Appui
    ax.add_patch(Rectangle((-0.4, -beam_h - 0.35), 0.8, 0.25,
                           facecolor="#333", edgecolor="#111", lw=1.0))
    ax.add_patch(FancyBboxPatch((-1.5, -beam_h - 0.95), 3.0, 0.55,
                                boxstyle="square,pad=0", facecolor=CONCRETE,
                                edgecolor=GREEN, lw=1.3))
    ax.text(0, -beam_h - 0.22, "Appui élastomère", fontsize=8, ha="center",
            color="white", fontweight="bold", va="center")
    ax.text(0, -beam_h - 0.7, "Chevêtre BA", fontsize=9, ha="center",
            color=GREEN, fontweight="bold")

    ax.set_xlim(-2.2, 3.2)
    ax.set_ylim(-beam_h - 1.4, 1.1)
    ax.set_title("VARIANTE B — Détail de la poutre précontrainte\n"
                 "et appui sur chevêtre",
                 fontsize=12, fontweight="bold", color=GREEN, pad=10)
    fig.tight_layout()
    path = OUT / "08_precontraint_detail_poutre.png"
    fig.savefig(path, dpi=220, bbox_inches="tight", facecolor=fig.get_facecolor())
    plt.close(fig)
    return path


def fig_comparatif():
    fig, axes = plt.subplots(1, 2, figsize=(16, 7), dpi=200)
    fig.patch.set_facecolor("#FAFAF7")

    # Coupes côté à côté
    for ax, kind in zip(axes, ["mixte", "pc"]):
        ax.set_facecolor("#FAFAF7")
        set_style(ax)
        half = W_TOTAL / 2
        if kind == "mixte":
            color = NAVY
            title = "A — Tablier mixte"
            beam_h = 0.90
            ax.add_patch(Rectangle((-half, 0), W_TOTAL, 0.22, facecolor="#D6DDE4",
                                   edgecolor=color, lw=1.4))
            for xc in np.linspace(-W_CHAUSSEE / 2, W_CHAUSSEE / 2, 5):
                ax.add_patch(Rectangle((xc - 0.18, -0.05), 0.36, 0.05, facecolor=STEEL, zorder=5))
                ax.add_patch(Rectangle((xc - 0.015, -beam_h + 0.05), 0.03, beam_h - 0.1,
                                       facecolor=STEEL_FILL, zorder=5))
                ax.add_patch(Rectangle((xc - 0.18, -beam_h), 0.36, 0.05, facecolor=STEEL, zorder=5))
            points = [
                "• Dalle BA + poutres acier I",
                "• Connexion par goujons",
                "• Tablier plus léger",
                "• Montage métallique rapide",
                "• Bon pour portées moyennes",
                f"• L = {L_TOTAL:.2f} m (4×{L_SPAN:.2f})",
            ]
        else:
            color = GREEN
            title = "B — Béton précontraint"
            beam_h = 1.10
            ax.add_patch(Rectangle((-half, 0), W_TOTAL, 0.20, facecolor="#D6DDE4",
                                   edgecolor=color, lw=1.4))
            for xc in np.linspace(-W_CHAUSSEE / 2, W_CHAUSSEE / 2, 5):
                top_w, bot_w, web_w = 0.40, 0.52, 0.12
                top_t, bot_t = 0.12, 0.15
                ax.add_patch(Polygon([
                    (xc - top_w / 2, 0), (xc + top_w / 2, 0),
                    (xc + top_w / 2, -top_t), (xc + web_w / 2, -top_t),
                    (xc + web_w / 2, -beam_h + bot_t), (xc + bot_w / 2, -beam_h + bot_t),
                    (xc + bot_w / 2, -beam_h), (xc - bot_w / 2, -beam_h),
                    (xc - bot_w / 2, -beam_h + bot_t), (xc - web_w / 2, -beam_h + bot_t),
                    (xc - web_w / 2, -top_t), (xc - top_w / 2, -top_t),
                ], closed=True, facecolor=PC_BEAM, edgecolor=color, lw=1.1))
                ax.add_patch(Circle((xc, -beam_h + 0.35), 0.04, facecolor=ACCENT, zorder=6))
            points = [
                "• Poutres I préfabriquées",
                "• Précontrainte par torons",
                "• Qualité usine maîtrisée",
                "• Grande durabilité",
                "• Entretien réduit",
                f"• L = {L_TOTAL:.2f} m (4×{L_SPAN:.2f})",
            ]

        ax.set_xlim(-half - 0.8, half + 0.8)
        ax.set_title(title, fontsize=12, fontweight="bold", color=color, pad=8)
        ax.text(0, -beam_h - 0.45, "\n".join(points), fontsize=8.5, ha="center",
                va="top", color="#333",
                bbox=dict(boxstyle="round,pad=0.4", facecolor="white",
                          edgecolor=color, lw=1.0))
        ax.set_ylim(-beam_h - 3.2, 0.9)

    fig.suptitle("COMPARAISON ARCHITECTURALE DES DEUX VARIANTES\n"
                 f"Pont — Portée totale {L_TOTAL:.2f} m  |  Largeur {W_TOTAL:.2f} m",
                 fontsize=14, fontweight="bold", color=NAVY, y=0.98)
    fig.tight_layout(rect=[0, 0, 1, 0.93])
    path = OUT / "09_comparatif_variantes.png"
    fig.savefig(path, dpi=220, bbox_inches="tight", facecolor=fig.get_facecolor())
    plt.close(fig)
    return path


def build_pptx(paths):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    NAVY_C = RGBColor(10, 37, 64)
    GREEN_C = RGBColor(47, 107, 79)
    WHITE = RGBColor(255, 255, 255)
    MUTED = RGBColor(90, 90, 90)
    ACC = RGBColor(196, 92, 38)

    def blank():
        return prs.slides.add_slide(prs.slide_layouts[6])

    def header(slide, title, color=NAVY_C):
        bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0),
                                     prs.slide_width, Inches(0.75))
        bar.fill.solid()
        bar.fill.fore_color.rgb = color
        bar.line.fill.background()
        tf = slide.shapes.add_textbox(Inches(0.4), Inches(0.15), Inches(12.5), Inches(0.5)).text_frame
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(20)
        p.font.bold = True
        p.font.color.rgb = WHITE

    def add_img(slide, path, left, top, width):
        slide.shapes.add_picture(str(path), Inches(left), Inches(top), width=Inches(width))

    # Cover
    s = blank()
    bg = s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0),
                            prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(247, 245, 241)
    bg.line.fill.background()
    header(s, "CONCEPTIONS ARCHITECTURALES — VARIANTES DE PONT")
    box = s.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.8), Inches(5)).text_frame
    lines = [
        ("Portée étudiée : 4 × 15,70 m = 62,80 m", 28, True, NAVY_C),
        ("", 14, False, MUTED),
        ("Variante A — Pont à tablier mixte (béton armé + poutres acier)", 20, True, NAVY_C),
        ("Variante B — Pont en béton précontraint (poutres préfabriquées)", 20, True, GREEN_C),
        ("", 14, False, MUTED),
        ("Vues : profil longitudinal • vue en plan • coupe transversale • détail technique", 16, False, MUTED),
        ("", 14, False, MUTED),
        ("Réalisé par : Mamady KABA & Thierno Ibrahima Hassanatou Diallo", 16, False, MUTED),
        ("Encadré par : Monsieur Fodé Bangaly Keita", 16, False, MUTED),
        ("Mémoire d'ingénieurs — Ponts et chaussées", 16, False, MUTED),
    ]
    first = True
    for text, size, bold, color in lines:
        p = box.paragraphs[0] if first else box.add_paragraph()
        first = False
        p.text = text
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = color

    # Mixte slides
    titles_mixte = [
        ("VARIANTE A — Profil longitudinal (tablier mixte)", paths[0], NAVY_C),
        ("VARIANTE A — Vue en plan", paths[1], NAVY_C),
        ("VARIANTE A — Coupe transversale", paths[2], NAVY_C),
        ("VARIANTE A — Détail de connexion mixte", paths[3], NAVY_C),
    ]
    for title, path, color in titles_mixte:
        s = blank()
        header(s, title, color)
        add_img(s, path, 0.35, 0.95, 12.6)

    # PC slides
    titles_pc = [
        ("VARIANTE B — Profil longitudinal (béton précontraint)", paths[4], GREEN_C),
        ("VARIANTE B — Vue en plan", paths[5], GREEN_C),
        ("VARIANTE B — Coupe transversale", paths[6], GREEN_C),
        ("VARIANTE B — Détail de poutre précontrainte", paths[7], GREEN_C),
    ]
    for title, path, color in titles_pc:
        s = blank()
        header(s, title, color)
        add_img(s, path, 0.35, 0.95, 12.6)

    # Comparatif
    s = blank()
    header(s, "COMPARAISON DES DEUX VARIANTES ARCHITECTURALES", ACC)
    add_img(s, paths[8], 0.35, 0.95, 12.6)

    # Synthèse
    s = blank()
    header(s, "SYNTHÈSE CONCEPTUELLE POUR LE MÉMOIRE")
    tf = s.shapes.add_textbox(Inches(0.7), Inches(1.2), Inches(12), Inches(5.5)).text_frame
    bullets = [
        "Les deux variantes respectent la même géométrie : 4 travées de 15,70 m (L = 62,80 m), largeur utile 8,00 m.",
        "Variante A (mixte) : dalle BA collaborante sur poutres acier — tablier léger, montage rapide, hauteur constructive réduite.",
        "Variante B (précontraint) : poutres I préfabriquées + hourdis — qualité usine, durabilité, câbles en parabole.",
        "Appuis communs : culées H = 5 m, 3 piles, chevêtres, appareils d'appui élastomères, semelles.",
        "Hydraulique commune : NPBE = 1,00 m • NPHE = 4,00 m • Affouillement prévu = 5,00 m.",
        "La variante BA (déjà retenue dans le mémoire) reste la référence ; A et B servent à la comparaison multicritère.",
    ]
    first = True
    for b in bullets:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.text = "•  " + b
        p.font.size = Pt(17)
        p.font.color.rgb = RGBColor(30, 30, 30)
        p.space_after = Pt(10)

    out = Path("/workspace/exports/Conceptions_Architecturales_Variantes_Pont.pptx")
    prs.save(out)
    return out


def build_pdf(paths):
    """Assembler un PDF multi-pages à partir des PNG."""
    images = [PILImage.open(p).convert("RGB") for p in paths]
    out = Path("/workspace/exports/Conceptions_Architecturales_Variantes_Pont.pdf")
    images[0].save(out, save_all=True, append_images=images[1:], resolution=220)
    return out


def main():
    print("Génération des vues variante A (mixte)...")
    p1 = fig_mixte_longitudinal()
    p2 = fig_mixte_plan()
    p3 = fig_mixte_coupe()
    p4 = fig_mixte_detail()
    print("Génération des vues variante B (précontraint)...")
    p5 = fig_pc_longitudinal()
    p6 = fig_pc_plan()
    p7 = fig_pc_coupe()
    p8 = fig_pc_detail()
    print("Génération du comparatif...")
    p9 = fig_comparatif()
    paths = [p1, p2, p3, p4, p5, p6, p7, p8, p9]
    print("Assemblage PPTX + PDF...")
    pptx = build_pptx(paths)
    pdf = build_pdf(paths)
    print("OK")
    for p in paths:
        print(p)
    print(pptx)
    print(pdf)


if __name__ == "__main__":
    main()
